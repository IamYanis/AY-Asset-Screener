"""Synthetic format checks plus source-local numerical/provenance validation.

Synthetic values exist only in temporary test files, never in production outputs.
Source checks skip when the separately downloaded, gitignored inputs are absent.
"""
import csv
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

import pandas as pd
for dependency in ("openpyxl", "pdfplumber"):
    if importlib.util.find_spec(dependency) is None:
        raise unittest.SkipTest("Install research/requirements.txt to run institutional ingestion tests")
from openpyxl import Workbook, load_workbook
from research.ingest import (ROOT, SECTORS, checked_status, digest, imf_tables,
    jpmorgan_tables, number, oecd_tables, read_csv, run, schwab_tables, vanguard_tables)


def make_pdf(path, pages):
    """Small real text PDFs for parser tests, without a PDF-writer dependency."""
    font_id = 3 + 2 * len(pages)
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               ("<< /Type /Pages /Kids [" + " ".join(f"{3+2*i} 0 R" for i in range(len(pages))) + f"] /Count {len(pages)} >>").encode()]
    for i, lines in enumerate(pages):
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 900 1200] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {4+2*i} 0 R >>".encode())
        commands = ["BT /F1 10 Tf 30 1150 Td 12 TL"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({escaped}) Tj T*")
        stream = ("\n".join(commands) + "\nET").encode("cp1252")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    pdf, offsets = bytearray(b"%PDF-1.4\n"), [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(pdf)


class FormatTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_csv_bom_delimiter_multiline_and_escaped_quotes(self):
        path = self.root / "test.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=";")
            writer.writerow(["id", "note", "value"])
            writer.writerow(["001", 'A "quoted", multiline\nnote', "0"])
        frame, info = read_csv(path)
        self.assertEqual(info["encoding"], "utf-8-sig")
        self.assertEqual(info["delimiter"], ";")
        self.assertEqual(info["data_rows"], 1)
        self.assertEqual(frame.id.iloc[0], "001")
        self.assertIn('"quoted", multiline\nnote', frame.note.iloc[0])
        path.write_text("a,b\n1,2,3\n")
        with self.assertRaises((ValueError, csv.Error)):
            read_csv(path)

    def test_missing_invalid_negative_and_zero(self):
        for value in [None, "", "..", "n/a", "--"]:
            self.assertEqual(number(value), (None, "missing"))
        for value in ["not a forecast", "inf", float("inf")]:
            self.assertEqual(number(value), (None, "invalid_numeric"))
        self.assertEqual(number("0"), (0, "numeric"))
        self.assertEqual(number("-2.5"), (-2.5, "numeric"))

    def test_imf_wide_schema_cutoffs_and_missing_cells(self):
        frame = pd.DataFrame([dict(SERIES_CODE="AAA.GDP.A", COUNTRY="Testland", INDICATOR="GDP growth",
            UNIT="Percent", SCALE="Units", FREQUENCY="Annual", PUBLICATION_DATE="2026-10-13T02:00:00Z",
            LATEST_ACTUAL_ANNUAL_DATA="2025", PRIMARY_DOMESTIC_CURRENCY="Test dollars",
            **{"2025": "0", "2026": "-1.2", "2027": ".."})])
        path = self.root / "imf.csv"
        frame.to_csv(path, index=False, encoding="utf-8-sig")
        tables, _ = imf_tables(path)
        s, obs = tables["imf_series"], tables["imf_observations"].set_index("period")
        self.assertEqual(s.country_code.iloc[0], "AAA")
        self.assertEqual(s.currency.iloc[0], "")  # No currency for a percent rate.
        self.assertEqual(obs.loc[2025, "observation_kind"], "historical_actual")
        self.assertEqual(obs.loc[2026, "observation_kind"], "estimate_or_projection")
        self.assertEqual(obs.loc[2026, "value"], -1.2)
        self.assertTrue(pd.isna(obs.loc[2027, "value"]))
        self.assertEqual(obs.loc[2027, "value_status"], "missing")
        frame["LATEST_ACTUAL_ANNUAL_DATA"] = "FY2025/26"
        frame.to_csv(path, index=False)
        tables, _ = imf_tables(path)
        self.assertEqual(tables["imf_observations"].iloc[0].observation_kind, "historical_status_unknown")

    def test_oecd_dimensions_periods_and_blank_status(self):
        frame = pd.DataFrame([dict(REF_AREA="AAA", **{"Reference area": "Testland"}, MEASURE="GDPV_ANNPCT",
            Measure="Real GDP growth", FREQ="A", TIME_PERIOD=str(year), OBS_VALUE=value,
            OBS_STATUS="", UNIT_MEASURE="PA", UNIT_MULT="0", CURRENCY="", STRUCTURE_NAME="Economic Outlook 119")
            for year, value in [(2025, "0"), (2026, "-2"), (2027, "..")]])
        path = self.root / "oecd.csv"
        frame.to_csv(path, index=False)
        tables, _ = oecd_tables(path, "2026-06-03")
        obs = tables["oecd_observations"]
        self.assertEqual(obs.observation_kind.tolist(), ["historical_status_unknown", "current_year_estimate_or_forecast", "forecast"])
        self.assertEqual(obs.value.iloc[0], 0)
        self.assertTrue(pd.isna(obs.value.iloc[-1]))
        self.assertTrue(obs.CURRENCY.eq("").all())
        self.assertTrue(obs.OBS_STATUS.eq("").all())
        pd.concat([frame, frame.iloc[[0]]]).to_csv(path, index=False)
        with self.assertRaisesRegex(ValueError, "duplicate"):
            oecd_tables(path, "2026-06-03")

    def test_excel_column_mapping_currency_and_missing_value(self):
        workbook = Workbook()
        ws = workbook.active
        ws.title = "2026"
        ws["A2"] = "2026 Long-Term Capital Market Assumptions – USD"
        for cell, text in {"B5": "Compound Return 2025 (%)", "B6": "Annualized Volatility (%)",
                           "B7": "Arithmetic Return 2026 (%)", "B8": "Compound Return 2026 (%)"}.items():
            ws[cell] = text
        ws["G5"] = ws["B9"] = "U.S. Large Cap"
        ws["A9"] = "EQUITIES"
        ws["C9"], ws["D9"], ws["F9"] = 6.7, 7.94, 99  # Prior vintage must not leak.
        path = self.root / "matrix.xlsx"
        workbook.save(path)
        metadata = {"publication_date": "2025-10-20", "as_of_date": "2025-09-30"}
        tables, info = jpmorgan_tables(path, metadata)
        result = tables["jpmorgan_assumptions"].set_index("metric")
        self.assertEqual(result.loc["compound_return", "value"], 6.7)
        self.assertEqual(result.loc["arithmetic_return", "source_locator"], "'2026'!D9")
        self.assertTrue(pd.isna(result.loc["annualised_volatility", "value"]))
        self.assertEqual(result.loc["annualised_volatility", "value_status"], "missing")
        self.assertTrue(result.currency.eq("USD").all())
        self.assertTrue(result.horizon_min_years.eq(10).all())
        self.assertEqual(info["sheet_names"], ["2026"])
        ws["A2"] = "2026 Long-Term Capital Market Assumptions – EUR"
        workbook.save(path)
        with self.assertRaisesRegex(ValueError, "USD"):
            jpmorgan_tables(path, metadata)
        workbook.close()

    def test_pdf_vanguard_narrative_ranges_review_and_missing_range(self):
        path = self.root / "vanguard.pdf"
        pages = [["Vanguard Capital Markets Model forecasts", "July 22, 2026"], [
            "U.S. equities: Our 10-year expected annualized return declined from a range of 4.9%–6.9% to a range of 4.2%–6.2%.",
            "International equities: Expected returns for developed markets ex-U.S. fell from a range of 5.4%–7.4% to a range of 4.5%-6.5%,",
            "while emerging markets expected returns fell from a range of 3.6%–5.6% to a range of 2%–4%.",
            "Outlooks are based on a June 30, 2026, running of the model."],
            ["Returns are in nominal terms, in U.S. dollar terms.", "Chart with no numerical table."]]
        make_pdf(path, pages)
        tables, info = vanguard_tables(path, {})
        ranges = tables["vanguard_return_ranges"]
        self.assertEqual(ranges.lower_pct.tolist(), [4.2, 4.5, 2])
        self.assertEqual(ranges.upper_pct.tolist(), [6.2, 6.5, 4])
        self.assertEqual(info["publication_date"], "2026-07-22")
        self.assertEqual(info["as_of_date"], "2026-06-30")
        self.assertTrue(ranges.extraction_status.eq("unverified_pdf").all())
        review = {"vanguard": {"sha256": digest(path), "ranges": {
            r.asset_class: [r.lower_pct, r.upper_pct] for r in ranges.itertuples()}}}
        checked, _ = vanguard_tables(path, review)
        self.assertTrue(checked["vanguard_return_ranges"].extraction_status.eq("checked_against_local_pdf").all())
        pages[1][2] = "Emerging markets numerical range not captured."
        make_pdf(path, pages)
        incomplete, info = vanguard_tables(path, review)
        self.assertEqual(len(incomplete["vanguard_return_ranges"]), 2)
        self.assertEqual(info["extraction_status"], "unverified_pdf")
        self.assertTrue(any("No exact narrative range" in issue for issue in info["issues"]))
        # Removing the US range must not misattribute the developed-market range.
        pages[1][0] = "U.S. equities: Our 10-year expected annualized return is not printed."
        make_pdf(path, pages)
        incomplete, _ = vanguard_tables(path, review)
        self.assertEqual(incomplete["vanguard_return_ranges"].asset_class.tolist(),
                         ["Developed markets ex-U.S. equities (unhedged)"])
        pages[1][1] = "International equities: No developed-market numerical range printed."
        make_pdf(path, pages)
        empty, info = vanguard_tables(path, review)
        self.assertTrue(empty["vanguard_return_ranges"].empty)
        self.assertIn("lower_pct", empty["vanguard_return_ranges"].columns)
        self.assertEqual(len(info["issues"]), 4)

    def test_missing_input_does_not_create_outputs(self):
        with self.assertRaisesRegex(ValueError, "expected one input"):
            run(input_dir=self.root / "missing", output_dir=self.root / "out")
        self.assertFalse((self.root / "out").exists())

    def test_pdf_schwab_qualitative_evidence_and_disagreement(self):
        path = self.root / "schwab.pdf"
        pages = [["October 2, 2026", "Our six- to 12-month S&P 500 outlook.", "Sector ratings"] +
                 [f"{sector.replace(' ', '')} Neutral" for sector in SECTORS] + ["Legend:"], []]
        for sector in SECTORS:
            pages[1] += [f"{sector} sector (rating: Neutral)", "Positives: Test support.", "Risks: Test risk."]
        pages[1] += ["How should I use Schwab Sector Views?"]
        make_pdf(path, pages)
        tables, info = schwab_tables(path, {})
        ratings = tables["schwab_sector_outlook"]
        self.assertEqual(len(ratings), 11)
        self.assertTrue(ratings.rating.eq("Neutral").all())
        self.assertTrue(ratings.currency.eq("").all())
        self.assertTrue(ratings.positive_evidence.eq("Test support.").all())
        self.assertEqual(info["extraction_status"], "unverified_pdf")
        pages[0][3] = "CommunicationServices Morefavored"
        make_pdf(path, pages)
        with self.assertRaisesRegex(ValueError, "disagree"):
            schwab_tables(path, {})


class SourceValidationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        metadata = json.loads((ROOT / "source_metadata.json").read_text())
        if any(len(list((ROOT / "institutional").glob(m["pattern"]))) != 1 for m in metadata.values()):
            raise unittest.SkipTest("Five local licensed source files are not present; synthetic ingestion tests still run")
        cls.temp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temp.cleanup)
        cls.out = Path(cls.temp.name) / "processed"
        cls.tables, cls.inventory, cls.summary = run(output_dir=cls.out)

    def test_actual_imf_numbers_and_complete_cell_alignment(self):
        source = ROOT / "institutional" / self.inventory["imf"]["relative_path"]
        raw = pd.read_csv(source, encoding="utf-8-sig", dtype=str, keep_default_na=False).set_index("SERIES_CODE")
        output = self.tables["imf_observations"]
        # Compare every raw year cell independently via a wide pivot, including blanks.
        pivot = output.pivot(index="SERIES_CODE", columns="source_column", values="raw_value")
        years = [c for c in raw if c.isdigit()]
        pd.testing.assert_frame_equal(pivot.loc[raw.index, years], raw[years], check_names=False)
        data = output.set_index(["SERIES_CODE", "period"])
        for series, year, expected in [("USA.NGDP_RPCH.A", 2026, 2.328), ("GBR.NGDP_RPCH.A", 2027, 1.222),
                                       ("CHN.NGDP_RPCH.A", 2031, 3.291), ("USA.PCPIPCH.A", 2026, 3.648)]:
            self.assertEqual(data.loc[(series, year), "value"], expected)
        self.assertEqual(output.value.notna().sum(), 361508)
        self.assertEqual(output.value.isna().sum(), 64996)
        self.assertEqual(self.inventory["imf"]["indicator_count"], 145)
        series = self.tables["imf_series"]
        self.assertTrue(series.loc[series.UNIT.eq("Percent"), "currency"].eq("").all())

    def test_actual_oecd_numerical_alignment_and_status(self):
        source = ROOT / "institutional" / self.inventory["oecd"]["relative_path"]
        raw = pd.read_csv(source, dtype=str, keep_default_na=False)
        output = self.tables["oecd_observations"]
        pd.testing.assert_series_equal(output.value, pd.to_numeric(raw.OBS_VALUE), check_names=False)
        self.assertEqual(output.MEASURE.unique().tolist(), ["GDPV_ANNPCT"])
        self.assertEqual(output.REF_AREA.nunique(), 58)
        self.assertEqual(output.loc[output.observation_kind.eq("forecast"), "period"].unique().tolist(), [2027])
        self.assertTrue(output.OBS_STATUS.eq("").all())

    def test_actual_excel_all_177_cells_match_and_percent_scale(self):
        source = ROOT / "institutional" / self.inventory["jpmorgan"]["relative_path"]
        workbook = load_workbook(source, data_only=True)
        self.addCleanup(workbook.close)
        output = self.tables["jpmorgan_assumptions"]
        for row in output.itertuples():
            sheet, address = row.source_locator.split("!")
            cell = workbook[sheet.strip("'")][address]
            self.assertEqual(row.value, cell.value)
            self.assertEqual(row.asset_class, workbook["2026"].cell(cell.row, 2).value)
        self.assertEqual(len(output), 177)
        large_cap = output[output.asset_class.eq("U.S. Large Cap")].set_index("metric")
        self.assertEqual(large_cap.loc["compound_return", "value"], 6.7)
        self.assertTrue(output.currency.eq("USD").all())
        self.assertTrue(output.horizon_max_years.eq(15).all())

    def test_actual_pdf_review_values_and_provenance(self):
        review = json.loads((ROOT / "pdf_review.json").read_text())
        schwab = self.tables["schwab_sector_outlook"]
        self.assertEqual(schwab.set_index("sector").rating.to_dict(), review["schwab"]["ratings"])
        self.assertTrue(schwab.extraction_status.eq("checked_against_local_pdf").all())
        self.assertTrue(schwab.positive_evidence.str.len().gt(20).all())
        self.assertTrue(schwab.risk_evidence.str.len().gt(20).all())
        vanguard = self.tables["vanguard_return_ranges"]
        self.assertEqual({r.asset_class: [r.lower_pct, r.upper_pct] for r in vanguard.itertuples()}, review["vanguard"]["ranges"])
        self.assertTrue(vanguard.extraction_status.eq("checked_against_local_pdf").all())
        self.assertEqual(set(vanguard.publication_date), {"2026-07-22"})
        self.assertEqual(set(vanguard.as_of_date), {"2026-06-30"})
        self.assertTrue(vanguard.horizon_max_years.eq(10).all())
        for name, table in self.tables.items():
            self.assertTrue(table.source_id.isin(self.inventory).all(), name)
        for key, entry in self.inventory.items():
            path = ROOT / "institutional" / entry["relative_path"]
            self.assertEqual(entry["sha256"], digest(path))
            self.assertEqual(entry["filename"], path.name)
            self.assertTrue(entry["url"].startswith("https://"))
        self.assertEqual(set(json.loads((self.out / "inventory.json").read_text())), set(self.inventory))
        self.assertTrue((self.out.parent / "data_quality_report.md").exists())

    def test_repeat_run_is_byte_reproducible(self):
        repeated = Path(self.temp.name) / "repeat" / "processed"
        run(output_dir=repeated)
        for path in self.out.iterdir():
            self.assertEqual(digest(path), digest(repeated / path.name), path.name)
        self.assertEqual(digest(self.out.parent / "data_quality_report.md"),
                         digest(repeated.parent / "data_quality_report.md"))


if __name__ == "__main__":
    unittest.main()

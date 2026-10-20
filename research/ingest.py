"""Offline Stage 1 ingestion. Run: python -m research.ingest

Inputs are never downloaded or modified. Source-specific tables stay separate;
blank/invalid values stay missing. See research/README.md for the output joins.
"""
import argparse
import csv
import hashlib
import io
import json
import math
import re
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SECTORS = ["Communication Services", "Consumer Discretionary", "Consumer Staples",
           "Energy", "Financials", "Health Care", "Industrials",
           "Information Technology", "Materials", "Real Estate", "Utilities"]
RATINGS = ["Least Favored", "Less Favored", "Neutral", "More Favored", "Most Favored"]
MISSING = {"", "--", "..", "...", "n/a", "na", "nan", "nd", "—"}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_csv(path):
    """Inspect before parsing; strict UTF-8, actual delimiter, logical CSV rows."""
    raw = Path(path).read_bytes()
    encoding = "utf-8-sig" if raw.startswith(b"\xef\xbb\xbf") else "utf-8"
    text = raw.decode(encoding)  # No lossy decoding or guessed legacy encoding.
    dialect = csv.Sniffer().sniff(text[:32000], delimiters=",;\t")
    # Sniff only the delimiter: a short sample can miss escaped quotes later in IMF notes.
    rows = list(csv.reader(io.StringIO(text), delimiter=dialect.delimiter, doublequote=True))
    if not rows or len(rows[0]) != len(set(rows[0])):
        raise ValueError(f"{path}: empty file or duplicate CSV headers")
    if any(len(row) != len(rows[0]) for row in rows):
        raise ValueError(f"{path}: inconsistent CSV field counts")
    frame = pd.DataFrame(rows[1:], columns=rows[0])
    return frame, {"encoding": encoding, "delimiter": dialect.delimiter,
                   "data_rows": len(frame), "columns": list(frame.columns),
                   "column_count": len(frame.columns), "physical_lines": len(text.splitlines()), "format": "CSV"}


def require(frame, columns):
    if missing := set(columns) - set(frame):
        raise ValueError(f"Missing source columns: {sorted(missing)}")


def number(value):
    """Return (finite numeric value, status), preserving real zero and negatives."""
    token = "" if value is None else str(value).strip()
    if token.lower() in MISSING:
        return None, "missing"
    try:
        value = float(token)
        if math.isfinite(value):
            return value, "numeric"
    except (ValueError, TypeError):
        pass
    return None, "invalid_numeric"


def imf_tables(path):
    frame, info = read_csv(path)
    require(frame, ["SERIES_CODE", "COUNTRY", "INDICATOR", "UNIT", "SCALE",
                    "FREQUENCY", "PUBLICATION_DATE", "LATEST_ACTUAL_ANNUAL_DATA",
                    "PRIMARY_DOMESTIC_CURRENCY"])
    if not frame.FREQUENCY.eq("Annual").all() or frame.SERIES_CODE.duplicated().any():
        raise ValueError("IMF: expected unique annual series")
    years = [c for c in frame if re.fullmatch(r"\d{4}", c)]
    if not years:
        raise ValueError("IMF: no annual value columns")
    series = frame.drop(columns=years).copy()
    series.insert(0, "source_id", "imf")
    series.insert(1, "source_row", range(2, len(series) + 2))
    codes = series.SERIES_CODE.str.split(".")
    if not codes.map(len).eq(3).all():
        raise ValueError("IMF: expected COUNTRY.INDICATOR.FREQUENCY series codes")
    series["country_code"] = codes.str[0]
    series["metric_code"] = codes.str[1]
    # Currency is a property of the metric, not the country or a percent unit.
    def currency(row):
        label = row.INDICATOR.lower()
        if row.UNIT in {"Percent", "Index"}:
            return ""
        if "domestic currency per international dollar" in label:
            return "domestic currency per PPP international dollar"
        if "international dollar" in label:
            return "PPP international dollar"
        if "domestic currency" in label or row.UNIT == "Domestic currency":
            return row.PRIMARY_DOMESTIC_CURRENCY
        if "us dollar" in label or "us cent" in label or row.UNIT.startswith("US dollar") or row.UNIT.startswith("US cent"):
            return "USD"
        return ""
    series["currency"] = series.apply(currency, axis=1)
    obs = frame[["SERIES_CODE"] + years].melt(id_vars="SERIES_CODE", var_name="period", value_name="raw_value")
    obs = obs.merge(series[["SERIES_CODE", "source_row", "PUBLICATION_DATE", "LATEST_ACTUAL_ANNUAL_DATA"]],
                    on="SERIES_CODE", validate="many_to_one")
    obs["source_id"] = "imf"
    obs["source_column"] = obs.period
    parsed = obs.raw_value.map(number)
    obs["value"] = parsed.str[0]
    obs["value_status"] = parsed.str[1]
    obs["period"] = obs.period.astype(int)
    publication_year = pd.to_datetime(obs.PUBLICATION_DATE, errors="raise", utc=True).dt.year
    actual = pd.to_numeric(obs.LATEST_ACTUAL_ANNUAL_DATA.where(obs.LATEST_ACTUAL_ANNUAL_DATA.str.fullmatch(r"\d{4}")), errors="coerce")
    obs["observation_kind"] = "historical_status_unknown"
    obs.loc[actual.notna() & obs.period.le(actual), "observation_kind"] = "historical_actual"
    obs.loc[actual.notna() & obs.period.gt(actual), "observation_kind"] = "estimate_or_projection"
    obs.loc[obs.period.gt(publication_year), "observation_kind"] = "forecast"
    obs["years_from_publication"] = obs.period - publication_year
    obs["classification_basis"] = "latest_actual_if_plain_year; future_period_vs_publication"
    obs = obs.drop(columns=["PUBLICATION_DATE", "LATEST_ACTUAL_ANNUAL_DATA"])
    info.update({"period_min": min(years), "period_max": max(years),
                 "geographies": series[["country_code", "COUNTRY"]].drop_duplicates().to_dict("records"),
                 "indicator_count": series.metric_code.nunique(),
                 "publication_dates": sorted(series.PUBLICATION_DATE.unique()),
                 "non_year_actual_cutoffs": sorted(series.loc[~series.LATEST_ACTUAL_ANNUAL_DATA.str.fullmatch(r"\d{4}"), "LATEST_ACTUAL_ANNUAL_DATA"].unique())})
    return {"imf_series": series, "imf_observations": obs}, info


def oecd_tables(path, publication_date):
    frame, info = read_csv(path)
    require(frame, ["REF_AREA", "Reference area", "MEASURE", "Measure", "FREQ", "TIME_PERIOD",
                    "OBS_VALUE", "OBS_STATUS", "UNIT_MEASURE", "UNIT_MULT", "CURRENCY", "STRUCTURE_NAME"])
    if not frame.FREQ.eq("A").all() or set(frame.STRUCTURE_NAME) != {"Economic Outlook 119"}:
        raise ValueError("OECD: adapter expects annual Economic Outlook 119")
    if frame.duplicated(["REF_AREA", "MEASURE", "TIME_PERIOD"]).any():
        raise ValueError("OECD: duplicate series-period observations")
    frame.insert(0, "source_id", "oecd")
    frame.insert(1, "source_row", range(2, len(frame) + 2))
    frame["source_column"] = "OBS_VALUE"
    frame["publication_date"] = publication_date
    frame["period"] = pd.to_numeric(frame.TIME_PERIOD, errors="raise").astype(int)
    parsed = frame.OBS_VALUE.map(number)
    frame["value"] = parsed.str[0]
    frame["value_status"] = parsed.str[1]
    frame["observation_kind"] = "historical_status_unknown"
    year = pd.Timestamp(publication_date).year
    frame.loc[frame.period.eq(year), "observation_kind"] = "current_year_estimate_or_forecast"
    frame.loc[frame.period.gt(year), "observation_kind"] = "forecast"
    frame["classification_basis"] = "period_vs_external_publication; OBS_STATUS_blank_not_verified"
    frame["years_from_publication"] = frame.period - year
    info.update({"period_min": int(frame.period.min()), "period_max": int(frame.period.max()),
                 "geographies": frame[["REF_AREA", "Reference area"]].drop_duplicates().to_dict("records"),
                 "indicators": frame[["MEASURE", "Measure"]].drop_duplicates().to_dict("records"),
                 "observation_statuses": sorted(frame.OBS_STATUS.unique())})
    return {"oecd_observations": frame}, info


def pdf_pages(path):
    with pdfplumber.open(path) as pdf:
        return [p.extract_text() or "" for p in pdf.pages], {
            "format": "PDF", "page_count": len(pdf.pages), "pdf_title": pdf.metadata.get("Title", ""),
            "pdf_creation_date": pdf.metadata.get("CreationDate", "")}


def publication_date(text):
    match = re.search(r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December) \d{1,2}, \d{4}\b", text)
    if not match:
        raise ValueError("PDF: publication date missing from article byline")
    return pd.Timestamp(match[0]).date().isoformat()


def checked_status(path, institution, records, review):
    """A different file or changed values cannot inherit a previous PDF review."""
    entry = review.get(institution, {})
    if digest(path) != entry.get("sha256"):
        return "unverified_pdf"
    actual = ({r["sector"]: r["rating"] for r in records} if institution == "schwab" else
              {r["asset_class"]: [r["lower_pct"], r["upper_pct"]] for r in records})
    expected = entry.get("ratings" if institution == "schwab" else "ranges")
    return "checked_against_local_pdf" if actual == expected else "unverified_pdf"


def schwab_tables(path, review):
    pages, info = pdf_pages(path)
    date = publication_date(pages[0])
    text = "\n".join(pages)
    if "six- to 12-month" not in text or "S&P 500" not in text:
        raise ValueError("Schwab: missing horizon/benchmark evidence")
    table = text.split("Sector ratings", 1)[1].split("Legend:", 1)[0]
    table = re.sub(r"\s", "", table).lower()
    headers = list(re.finditer(r"(?m)^(.+?) sector \(rating: (.+?)\)\s*$", text))
    records = []
    for i, match in enumerate(headers):
        sector, rating = match.group(1), match.group(2)
        if sector not in SECTORS or rating not in RATINGS:
            raise ValueError("Schwab: unknown sector/rating in commentary")
        compact = re.sub(r"\s", "", sector).lower()
        table_ratings = {r for r in RATINGS if compact + r.replace(" ", "").lower() in table}
        if table_ratings != {rating}:
            raise ValueError(f"Schwab: rating table/commentary disagree for {sector}")
        section = text[match.end():headers[i + 1].start() if i + 1 < len(headers) else text.index("How should I use Schwab Sector Views?")]
        if "Positives:" not in section or "Risks:" not in section:
            raise ValueError(f"Schwab: missing supporting text for {sector}")
        positive, risk = section.split("Positives:", 1)[1].split("Risks:", 1)
        page = next(n for n, t in enumerate(pages, 1) if match.group(0).strip() in t)
        records.append({"source_id": "schwab", "publication_date": date, "as_of_date": date,
            "region": "United States", "sector": sector, "rating": rating,
            "metric": "relative_sector_outlook", "unit": "qualitative ordinal category",
            "benchmark": "S&P 500", "currency": "", "currency_status": "not_applicable_to_rating",
            "horizon_min_months": 6, "horizon_max_months": 12,
            "source_locator": f"rating table pp. 2-3; commentary p. {page}",
            "positive_evidence": " ".join(positive.split()), "risk_evidence": " ".join(risk.split()),
            "evidence_status": "extracted_prose_not_independently_validated"})
    if len(records) != 11 or {r["sector"] for r in records} != set(SECTORS):
        raise ValueError("Schwab: incomplete or duplicate sector set")
    status = checked_status(path, "schwab", records, review)
    for r in records:
        r["extraction_status"] = status
    info.update({"publication_date": date, "as_of_date": date, "extraction_status": status})
    return {"schwab_sector_outlook": pd.DataFrame(records)}, info


def vanguard_tables(path, review):
    pages, info = pdf_pages(path)
    date = publication_date(pages[0])
    text = " ".join(" ".join(pages).split())
    if "U.S. dollar terms" not in text or "nominal terms" not in text:
        raise ValueError("Vanguard: missing currency/nominal basis evidence")
    model = re.search(r"based on a (\w+ \d{1,2}, \d{4}), running", text)
    if not model or "10-year expected annualized return" not in text:
        raise ValueError("Vanguard: missing model date or narrative horizon")
    as_of = pd.Timestamp(model[1]).date().isoformat()
    band = r"(-?\d+(?:\.\d+)?)%[–-](-?\d+(?:\.\d+)?)%"
    sections = [
        ("U.S. equities", "United States", "U.S. equities:", "International equities:"),
        ("Developed markets ex-U.S. equities (unhedged)", "Developed markets ex-US", "developed markets ex-U.S. fell", "while emerging markets expected returns"),
        ("Emerging markets equities (unhedged)", "Emerging markets", "while emerging markets expected returns", "Fixed income:")]
    records, issues = [], []
    narrative = " ".join(pages[1].split())
    for asset, region, start, stop in sections:
        # Limit matches to the relevant clause so a missing range cannot borrow
        # the following asset's numbers.
        segment = narrative.split(start, 1)[1].split(stop, 1)[0] if start in narrative else ""
        match = re.search(r"to a range of " + band, segment)
        if not match:
            issues.append(f"No exact narrative range extracted for {asset}; no observation created")
            continue
        low, high = map(float, match.groups())
        if low > high:
            raise ValueError("Vanguard: reversed return range")
        records.append({"source_id": "vanguard", "asset_class": asset, "region": region,
            "publication_date": date, "as_of_date": as_of, "currency": "USD", "unit": "percent per annum",
            "metric": "annualised_nominal_total_return_range", "lower_pct": low, "upper_pct": high,
            "horizon_min_years": 10, "horizon_max_years": 10,
            "range_basis": "central band around 50th percentile; not a confidence interval",
            "source_locator": "p. 2 narrative; p. 3 nominal USD definition",
            "evidence": match[0]})
    status = checked_status(path, "vanguard", records, review)
    for r in records:
        r["extraction_status"] = status
    info.update({"publication_date": date, "as_of_date": as_of, "extraction_status": status,
                 "issues": issues + ["Chart labels extract, but chart-only endpoint values, volatility, tail percentiles and 30-year numerical projections are not reliably available. No graph digitisation or table reconstruction attempted."]})
    columns = ["source_id", "asset_class", "region", "publication_date", "as_of_date", "currency", "unit",
               "metric", "lower_pct", "upper_pct", "horizon_min_years", "horizon_max_years", "range_basis",
               "source_locator", "evidence", "extraction_status"]
    return {"vanguard_return_ranges": pd.DataFrame(records, columns=columns)}, info


def jpmorgan_tables(path, metadata):
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if "2026" not in workbook.sheetnames:
            raise ValueError("JPM: expected 2026 sheet; inspect new workbook before ingestion")
        ws = workbook["2026"]
        title = str(ws["A2"].value)
        if not re.search(r"2026 Long-Term Capital Market Assumptions\s*[–-]\s*USD", title):
            raise ValueError("JPM: 2026 edition/USD currency not confirmed by workbook title")
        for cell, label in {"B8": "Compound Return 2026 (%)", "B7": "Arithmetic Return 2026 (%)", "B6": "Annualized Volatility (%)", "B5": "Compound Return 2025 (%)"}.items():
            if ws[cell].value != label:
                raise ValueError(f"JPM: layout changed at {cell}")
        names = [ws.cell(5, c).value for c in range(7, ws.max_column + 1) if ws.cell(5, c).value]
        if not names or len(names) != len(set(names)):
            raise ValueError("JPM: missing/duplicate matrix asset headers")
        records, group = [], ""
        for row, name in enumerate(names, 9):
            if ws.cell(row, 2).value != name:
                raise ValueError(f"JPM: matrix header and row label disagree at B{row}")
            group = ws.cell(row, 1).value or group
            for col, metric in [(3, "compound_return"), (4, "arithmetic_return"), (5, "annualised_volatility")]:
                cell = ws.cell(row, col)
                value, status = number(cell.value)
                records.append({"source_id": "jpmorgan", "asset_class": name, "source_group": group,
                    "metric": metric, "value": value, "raw_value": cell.value, "value_status": status,
                    "unit": "percent per annum" if col != 5 else "annualised percent volatility",
                    "currency": "USD", "publication_date": metadata["publication_date"],
                    "as_of_date": metadata["as_of_date"], "assumption_vintage": 2026,
                    "horizon_min_years": 10, "horizon_max_years": 15,
                    "source_locator": f"'2026'!{cell.coordinate}", "source_row": row,
                    "extraction_status": "workbook_numeric_cell",
                    "horizon_date_basis": "external_official_release_and_matrix_report",
                    "return_basis": "nominal; alternatives median manager net of fees; see official report p. 83"})
        info = {"format": "XLSX", "sheet_names": workbook.sheetnames, "sheet": ws.title,
                "worksheet_max_row": ws.max_row, "worksheet_max_column": ws.max_column,
                "data_rows": len(names), "data_row_range": f"9:{8 + len(names)}",
                "metrics_columns": {"C": "Compound Return 2026 (%)", "D": "Arithmetic Return 2026 (%)", "E": "Annualized Volatility (%)"},
                "ignored": "F: prior 2025 assumptions; G:BM lower-triangular correlations; empty formatted tail",
                "title": title.strip(), "audience_notice": ws["A1"].value,
                "publication_date": metadata["publication_date"], "as_of_date": metadata["as_of_date"]}
    finally:
        workbook.close()
    return {"jpmorgan_assumptions": pd.DataFrame(records)}, info


def run(input_dir=ROOT / "institutional", output_dir=ROOT / "processed"):
    input_dir, output_dir = Path(input_dir), Path(output_dir)
    registry = json.loads((ROOT / "source_metadata.json").read_text())
    review = json.loads((ROOT / "pdf_review.json").read_text())
    tables, inventory = {}, {}
    for institution, metadata in registry.items():
        matches = sorted(input_dir.glob(metadata["pattern"]))
        if len(matches) != 1:
            raise ValueError(f"{institution}: expected one input matching {metadata['pattern']}; found {len(matches)}")
        path = matches[0]
        if institution == "imf":
            extracted, info = imf_tables(path)
        elif institution == "oecd":
            extracted, info = oecd_tables(path, metadata["publication_date"])
        elif institution == "schwab":
            extracted, info = schwab_tables(path, review)
        elif institution == "vanguard":
            extracted, info = vanguard_tables(path, review)
        else:
            extracted, info = jpmorgan_tables(path, metadata)
        inventory[institution] = {**metadata, "filename": path.name, "relative_path": str(path.relative_to(input_dir)),
                                  "sha256": digest(path), "bytes": path.stat().st_size, **info}
        tables.update(extracted)
    # Parse every source before replacing output tables; never fill failed inputs.
    output_dir.mkdir(parents=True, exist_ok=True)
    for name, frame in tables.items():
        frame.to_csv(output_dir / f"{name}.csv", index=False, lineterminator="\n")
    summary = {}
    for name, frame in tables.items():
        summary[name] = {"rows": len(frame)}
        if "value_status" in frame:
            summary[name]["value_status_counts"] = frame.value_status.value_counts().to_dict()
        if "observation_kind" in frame:
            summary[name]["usable_by_kind"] = frame.loc[frame.value_status.eq("numeric"), "observation_kind"].value_counts().to_dict()
    (output_dir / "inventory.json").write_text(json.dumps(inventory, indent=2, ensure_ascii=False) + "\n")
    (output_dir / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    write_report(tables, inventory, summary, output_dir.parent / "data_quality_report.md")
    print(json.dumps(summary, indent=2))
    return tables, inventory, summary


def write_report(tables, inventory, summary, path):
    """Counts are generated from the same tables written to CSV; prose audits these snapshots."""
    entries = []
    for key, source in inventory.items():
        if "column_count" in source:
            shape = f"{source['data_rows']} × {source['column_count']} CSV; {source['encoding']}; delimiter {source['delimiter']!r}"
        elif "page_count" in source:
            shape = f"{source['page_count']} PDF pages"
        else:
            shape = f"sheet {source['sheet_names']}; {source['data_rows']} labelled rows; worksheet extent {source['worksheet_max_row']} × {source['worksheet_max_column']}"
        entries.append(f"- **{key}**: `{source['relative_path']}` — {source['bytes']:,} bytes; {shape}. SHA-256 `{source['sha256']}`.")
    inventory_lines = "\n".join(entries)
    counts = ["| Dataset | Output rows | Finite numeric / usable categories | Missing or invalid |",
              "|---|---:|---:|---:|"]
    for name, frame in tables.items():
        if name == "imf_series":
            continue
        usable = int(frame.value_status.eq("numeric").sum()) if "value_status" in frame else len(frame)
        counts.append(f"| {name} | {len(frame):,} | {usable:,} | {len(frame) - usable:,} |")
    counts_table = "\n".join(counts)
    imf_counts = json.dumps(summary["imf_observations"]["usable_by_kind"], sort_keys=True)
    oecd_counts = json.dumps(summary["oecd_observations"]["usable_by_kind"], sort_keys=True)
    imf = tables["imf_series"]
    imf_intro = (f"{len(imf):,} annual series; {imf.metric_code.nunique()} indicator codes/labels; "
                 f"{imf.country_code.nunique()} country/group identifiers. Annual columns 1980–2031. "
                 "Publication/update is 13 October 2026 (`2026-10-13T02:00:00Z`), not the filename's download time. "
                 "Countries include US, UK, euro-area economies, Japan, China, India and emerging economies; "
                 "groups include World, G7, EU, Euro Area, advanced/emerging markets and regional aggregates. "
                 "Commodity series are also present. Identifiers come from `COUNTRY.INDICATOR.A`; aggregate codes are not assumed ISO countries.")
    schwab_table = "\n".join(["|---|---|"] + [f"| {r.sector} | {r.rating} |" for r in tables["schwab_sector_outlook"].itertuples()])
    vanguard_table = "\n".join(["|---|---:|---:|---|"] + [f"| {r.asset_class} | {r.lower_pct:g} | {r.upper_pct:g} | 10 years |" for r in tables["vanguard_return_ranges"].itertuples()])
    schwab_status = inventory["schwab"]["extraction_status"]
    vanguard_status = inventory["vanguard"]["extraction_status"]
    report = f"""# Institutional research — Stage 1 data quality

Audit date: 15 October 2026. Offline extraction of the five requested files. No rankings, scores, recommendations, FX conversion or horizon blending.

## Inventory and provenance

{inventory_lines}

Every output joins by `source_id` to `processed/inventory.json` for the exact filename, SHA-256, publisher URL and metadata evidence. IMF observations additionally join by `SERIES_CODE` to `imf_series.csv` for indicators, units, scaling, currency and source publication date. Source row numbers count logical CSV records (header = 1); IMF quoted multiline metadata means they are not physical line numbers. PDF page numbers and Excel cell addresses are retained.

## Successfully extracted observations

{counts_table}

imf_observations usable classification counts: `{imf_counts}`.

oecd_observations usable classification counts: `{oecd_counts}`.

The numeric count includes historical observations; it is not a forecast count. Schwab rows are qualitative opinions, not numerical forecasts. Vanguard rows each contain a lower/upper range, not two independent forecasts. J.P. Morgan has three metrics per labelled assumption, including inflation (not an investable asset).

## IMF

{imf_intro}

Indicators cover nominal/real/PPP GDP and growth/per capita, CPI levels and average/end-period inflation, GDP deflator, unemployment/employment/population, investment/savings, fiscal revenue/expenditure/balances/debt, current/financial accounts, trade volumes, external debt/service, output gaps, and commodity prices/indices. Full indicator labels/codes and geography lists are in the metadata tables/inventory; not every indicator is available for every geography.

Units are source percent, index, domestic currency, USD, PPP international dollars and commodity-specific prices; scales are units, millions or billions. Values are preserved at source scale. No currency is assigned to unitless growth rates; domestic/PPP/USD metrics keep distinct currency descriptions. Empty source UNIT fields remain empty; metric text still describes per-capita/PPP meanings. Fiscal/calendar-year, base-year and methodology fields are retained.

`LATEST_ACTUAL_ANNUAL_DATA` is used only when it is a plain year. Post-cutoff values at/before publication year are estimates/projections, not confirmed future forecasts; 2027–2031 values are future-period projections by temporal inference. FY strings, blanks and 'Not applicable' are not silently converted to a calendar-year cutoff. Missing annual cells are kept with a missing status and never replaced with zero. Future-year coverage is incomplete and is counted only where a numeric value exists.

## OECD

The supplied filtered EO119 CSV has just **GDPV_ANNPCT: real GDP volume growth**, annual, percent per annum (`PA`, multiplier 0), 1961–2027, across 58 country/group identifiers. It does not contain the OECD's full inflation, employment or other forecast dataset. It includes US/UK/Japan/China/India, OECD countries and selected non-members, plus OECD, World, non-OECD, EA17 and other aggregates. EA17 is not interchangeable with IMF Euro Area membership.

Publication date is externally confirmed as 3 June 2026 for Economic Outlook 119 / 2026/1; the CSV contains no release date, data cutoff or populated OBS_STATUS. All currencies are blank, appropriate for a volume growth rate. 2027 is future-period forecast by date; 2026 is labelled current-year estimate-or-forecast using the release context, not an official observation status. Earlier rows retain unknown actual/estimate status. Base-year and all original dimensions are retained. Obtain EO119 last-historical-point metadata before claiming exact actual/forecast cutoffs per country.

## Schwab

Publication and ratings as-of date: **{inventory['schwab']['publication_date']}**, rather than the PDF print date. PDF extraction status: **{schwab_status}**. US S&P 500/GICS sectors only; six- to twelve-month relative outlook. Extracted ratings are cross-checked between the printed table and the commentary. Only `checked_against_local_pdf` means the current file hash and values also match the recorded visual review; `unverified_pdf` requires a fresh review. Supporting positive/risk prose is retained as extracted evidence, with page locators, and is not an independently validated forecast. Currency is not applicable to qualitative ratings.

| Sector | Qualitative rating |
{schwab_table}

Evidence covers AI/electrification investment and cyclical risks in industrials/materials, yield-curve support and credit risk in financials, innovation/regulatory risks in health care, and financing/valuation/consumer-demand pressures elsewhere. The local processed table preserves each sector's positive and negative rationale. Trailing six/twelve-month performance and concentration figures elsewhere in the PDF are historical context and are not extracted as forecasts. Categories remain text; no numeric favourability scores are assigned.

## Vanguard

Actual publication: **{inventory['vanguard']['publication_date']}** (22 July 2026 in the audited snapshot); model run/as-of: **{inventory['vanguard']['as_of_date']}**. The filename gives the model date. PDF extraction status: **{vanguard_status}**. The table contains only exact 10-year annualised **nominal total-return ranges, USD**, from the p. 2 narrative. The three ranges in the audited snapshot were visually checked; changed files/values require a fresh review. They are central bands around the model's median, not confidence intervals or a full outcome distribution. No midpoint is manufactured.

| Asset class | Lower (%) | Upper (%) | Horizon |
{vanguard_table}

The p. 3 chart contains additional equity, bond, commodity, inflation and currency labels and graphical ranges. Automatic table extraction yields labels and repeated axis text, **not numeric endpoint cells**. Visual inspection confirms a chart rather than a usable numeric table; the p. 2 horizon selector also visually highlights 30 years while the following chart says 10-year. Only explicitly 10-year narrative numbers are used. Fixed-income numbers, chart-only equities, median volatility, tail percentiles and all 30-year numeric forecasts remain unavailable. No pixel/axis estimates, OCR guesses or 10-to-30-year extrapolations are produced.

PDF review status is bound to SHA-256 and exact extracted ratings/ranges in `pdf_review.json`. A different file or different values returns `unverified_pdf`; it cannot inherit the review. Partial/missing narrative ranges produce issues instead of placeholder observations. Model outcomes remain hypothetical even when extraction is checked.

## J.P. Morgan

Actual workbook sheet: **2026**. A2 explicitly says **2026 Long-Term Capital Market Assumptions – USD**. Rows 9–67 hold 59 labelled assumptions, including inflation. Column C = 2026 compound return; D = 2026 arithmetic return; E = annualised volatility. Column F is prior-vintage 2025 compound return and is intentionally excluded; G:BM is the lower-triangular correlation matrix, also excluded. Merged header cells B5–B8 label these columns in reverse vertical order; blindly treating the first nonempty row as flat headers would be wrong. The 580 × 157 worksheet extent includes formatting and is not 580 observations.

177 finite numeric cells are extracted at original percentage-point scale; 6.7 means 6.7%, not 670%. Sample cross-check: U.S. Large Cap C31 = 6.7, D31 = 7.935086973158945, E31 = 16.47014248435096; UK Large Cap C37 = 6.6000000000000005. Workbook numbers are compared directly in tests. Printed matrix pp. 82–83 confirms column ordering and rounded values.

The **2026** label is an assumptions vintage, not a calendar-2026 return forecast. The official release dated **20 October 2025** specifies a **10–15-year** horizon; official matrix notes give **30 September 2025** as-of and nominal USD assumptions. These dates/horizon are external metadata, explicitly distinguished from workbook extraction. The workbook's 15 October 2025 modified timestamp is not its release date. A companion `ltcma-full-report.pdf` was also found locally during inspection and used only as a matrix cross-check; it is not a sixth ingestion input or required at runtime.

Coverage includes US, euro area, Japan, Hong Kong, UK, China, EAFE, emerging markets, Asia ex-Japan and global asset classes, plus real assets/private markets. Printed source groups are retained without inventing sector classifications. FX hedge/local-currency wording remains part of the exact asset label; USD investor-return currency does not mean every underlying security is USD-denominated. Alternatives are median-manager estimates net of fees; equity/bond assumptions are not active-manager alpha forecasts. A1 restricts this workbook to institutional/wholesale/professional/qualified investors and says not for retail use or distribution.

## Compatibility and sufficiency for a sector × region ranking

**Not sufficient for a defensible sector × region ranking.** Macro country GDP/inflation is not sector earnings or equity return. Schwab covers only US sectors, qualitatively over 6–12 months. Vanguard covers broad asset classes over 10 years with only three exact ranges recoverable here. J.P. Morgan covers broad asset classes over 10–15 years from a different vintage. None supplies a consistent non-US GICS sector panel. These data are therefore emitted separately, with no ranking, aggregation, sector imputation or horizon conversion.

Other incompatibilities: nominal vs real/PPP macro variables; rate vs level and millions/billions; domestic vs USD/PPP currency; FX hedge basis; arithmetic vs compound/geometric returns; net alternative-manager fees vs index assumptions; different release/as-of dates and fiscal calendars; nonmatching regional membership and aggregation weights; qualitative categories vs numerical ranges/point assumptions. No crosswalk equates similarly named aggregates.

A later heatmap needs a defined investment horizon and common as-of vintage, licensed sector-by-region earnings/revenue revisions, valuations, profitability and balance-sheet data, consistent GICS taxonomy and country/region membership, investable benchmark/asset identifiers and competition eligibility. It also needs comparable price/total-return/volatility data, FX/base-currency and hedging policy, and consistent sector weights, coverage and missingness. For outlook comparisons obtain non-US sector views at the same horizon, machine-readable Vanguard tables with volatility/percentiles, and OECD latest-actual metadata. Historical validation and explicit uncertainty/coverage rules would be needed before choosing any ranking method; no scoring is attempted in Stage 1.

## Extraction limitations and licensing

The pipeline is offline and schema-specific. It rejects changed CSV layouts, unknown sectors/ratings, inconsistent table/commentary, wrong workbook vintage/currency and missing inputs. Numeric tokens retain raw values and missing/invalid status. It neither executes Excel formulas nor reconstructs missing cached values. It does not infer chart numbers. Source URLs describe publisher provenance; original browser transfer URLs were not recorded. PDF print metadata can expose local download times and is not used as publication evidence. Re-running with changed content requires renewed PDF review and report assessment.

Downloaded files and **all derived data** in `research/processed/` are gitignored pending redistribution review. Code, metadata, tests and this limited audit report can be reviewed without publishing the source archive. Existing blank `research/sources.csv` is not treated as evidence. Supporting links were checked 15 October 2026; no broad republication rights are inferred from public download availability.

- [IMF statistical-data terms](https://www.imf.org/en/about/copyright-and-terms) allow attributed data reuse subject to their conditions; different terms apply to other content/third-party data and potential commercial use. Preserve IMF attribution and original metric meanings.
- [OECD terms, Data section](https://www.oecd.org/en/about/terms-conditions.html) permit attributed data reuse subject to dataset/third-party restrictions. Preserve the EO119 citation and verify any source-specific rights before redistribution.
- Schwab and Vanguard PDFs contain all-rights-reserved notices. Their availability is not a redistribution licence. Keep raw PDFs and extracted prose/tables local pending appropriate permission.
- J.P. Morgan's workbook audience/use restriction is explicit. Do not publish the workbook or derived matrix as a retail product; confirm permitted use and redistribution with the publisher before deployment.

## Date/horizon verification sources

- [IMF WEO source](https://data.imf.org/en/datasets/IMF.RES:WEO): exact release timestamp comes from the supplied CSV.
- [OECD Economic Outlook](https://www.oecd.org/en/topics/sub-issues/economic-outlook.html) and [EO119 explorer](https://data-explorer.oecd.org/?df[ag]=OECD.ECO.MAD&df[id]=DSD_EO@DF_EO&df[vs]=1.5): 3 June 2026 release, edition 2026/1.
- [Schwab sector outlook](https://www.schwab.com/learn/story/stock-sector-outlook): local PDF dated 2 October 2026.
- [Vanguard VCMM](https://corporate.vanguard.com/content/corporatesite/us/en/corp/vemo/vemo-return-forecasts.html): local PDF separates publication and model-run date.
- [J.P. Morgan 2026 release](https://am.jpmorgan.com/us/en/asset-management/adv/about-us/media/press-releases/jp-morgan-releases-2026-long-term-capital-market-assumptions/) and [official matrix report](https://am.jpmorgan.com/content/dam/jpm-am-aem/global/en/insights/portfolio-insights/ltcma/noindex/ltcma-full-report.pdf), pp. 82–83: release, horizon, as-of, currency and return conventions.
"""
    Path(path).write_text(report)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=ROOT / "institutional")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "processed")
    args = parser.parse_args()
    run(args.input_dir, args.output_dir)
